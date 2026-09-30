import io
import re
from pathlib import Path
import docx
import pypdf


def clean_extracted_text(text: str) -> str:
    if not text:
        return ""
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    no_controls = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", normalized)
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in no_controls.split("\n")]
    joined = "\n".join(lines)
    collapsed = re.sub(r"\n{3,}", "\n\n", joined)
    return collapsed.strip()


def extract_text_from_pdf(content: bytes) -> str:
    reader = pypdf.PdfReader(io.BytesIO(content))
    pages = []
    for page in reader.pages:
        extracted = page.extract_text()
        if extracted:
            pages.append(extracted)
    return clean_extracted_text("\n\n".join(pages))


def extract_text_from_docx(content: bytes) -> str:
    doc = docx.Document(io.BytesIO(content))
    chunks = []
    for paragraph in doc.paragraphs:
        txt = paragraph.text.strip()
        if txt:
            chunks.append(txt)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                chunks.append(" | ".join(cells))
    return clean_extracted_text("\n".join(chunks))


def extract_text_from_doc(content: bytes) -> str:
    if content.startswith(b"{\\rtf"):
        text = re.sub(r"\\(?:[a-zA-Z]+-?\d*|['\\][0-9a-fA-F]{2}|.)", "", content.decode("latin1", errors="ignore"))
        return clean_extracted_text(text)
    runs = re.findall(rb"[\x20-\x7E\t\r\n]{4,}", content)
    decoded = [r.decode("utf-8", errors="ignore").strip() for r in runs if r.strip()]
    return clean_extracted_text("\n".join(decoded))


def extract_resume_text(filename: str, content: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(content)
    elif ext == ".docx":
        return extract_text_from_docx(content)
    elif ext == ".doc":
        return extract_text_from_doc(content)
    raise ValueError(f"Unsupported file format: {ext}")
